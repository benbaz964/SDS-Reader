"""
template_filler.py
===================
Optional feature: user uploads a template (.docx or .txt/.md) containing
placeholders like {{product_name}}, {{cas_number_main}}, {{signal_word}},
{{composition_table}}, etc. matching the canonical variable names in
schema.py. We fill them in from an extracted SDSRecord.

.docx placeholders are matched even when Word has split "{{product_name}}"
across multiple runs (very common), by rebuilding paragraph text run-by-run.
"""

import re
from copy import deepcopy
from typing import Dict

from docx import Document
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")
ROW_PLACEHOLDER_RE = re.compile(r"\{\{\s*row:([a-zA-Z0-9_]+)\.([a-zA-Z0-9_]+)\s*\}\}")


def _stringify(value) -> str:
    if value is None:
        return "Not stated"
    if isinstance(value, list):
        if not value:
            return "Not stated"
        if isinstance(value[0], dict):
            # table-like field (e.g. composition_table) -> readable block
            lines = []
            for row in value:
                lines.append(", ".join(f"{k}: {v}" for k, v in row.items() if v))
            return "\n".join(lines)
        return ", ".join(str(v) for v in value)
    if value == "":
        return "Not stated"
    return str(value)


def fill_text_template(template_text: str, values: Dict) -> str:
    def repl(m):
        var = m.group(1)
        return _stringify(values.get(var, m.group(0)))
    return PLACEHOLDER_RE.sub(repl, template_text)


def _row_cell_field_map(row):
    """Returns {cell_index: (list_name, field_name)} for cells in this row
    that contain a row-placeholder, or {} if this isn't a template row."""
    mapping = {}
    for i, cell in enumerate(row.cells):
        text = "".join(p.text for p in cell.paragraphs)
        m = ROW_PLACEHOLDER_RE.search(text)
        if m:
            mapping[i] = (m.group(1), m.group(2))
    return mapping


def _set_cell_text(cell, text: str):
    # Replace the cell's first paragraph text (clears the rest) - simplest
    # reliable way to overwrite a templated cell.
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph(text)
        return
    first = paragraphs[0]
    for run in first.runs[1:]:
        run.text = ""
    if first.runs:
        first.runs[0].text = text
    else:
        first.add_run(text)
    for p in paragraphs[1:]:
        for run in p.runs:
            run.text = ""


def _clone_row(table: Table, template_row) -> "_Row":
    new_tr = deepcopy(template_row._tr)
    template_row._tr.addnext(new_tr)
    return _Row(new_tr, table)


def fill_dynamic_tables(doc: Document, values: Dict):
    """Finds any table row templated with {{row:listname.field}} and expands
    it into one row per item of values[listname] (a list of dicts, e.g. the
    per-ingredient substance_breakdown table)."""
    for table in doc.tables:
        template_row = None
        field_map = {}
        for row in table.rows:
            fm = _row_cell_field_map(row)
            if fm:
                template_row = row
                field_map = fm
                break
        if template_row is None:
            continue

        list_name = next(iter(field_map.values()))[0]
        items = values.get(list_name) or []

        if not items:
            for i, (lname, field) in field_map.items():
                _set_cell_text(template_row.cells[i], "None reported")
            continue

        last_row = template_row
        template_tr_copy = deepcopy(template_row._tr)
        for idx, item in enumerate(items):
            if idx == 0:
                target_row = template_row
            else:
                new_tr = deepcopy(template_tr_copy)
                last_row._tr.addnext(new_tr)
                target_row = _Row(new_tr, table)
                last_row = target_row
            for i, (lname, field) in field_map.items():
                _set_cell_text(target_row.cells[i], str(item.get(field, "")) or "Not stated")


def _fill_paragraph(paragraph: Paragraph, values: Dict):
    full_text = "".join(run.text for run in paragraph.runs)
    if "{{" not in full_text:
        return
    new_text = fill_text_template(full_text, values)
    if new_text == full_text:
        return
    # Wipe existing runs, write result into the first run to preserve some formatting.
    for run in paragraph.runs[1:]:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
    else:
        paragraph.add_run(new_text)


def _walk_tables(doc_or_cell, values: Dict):
    for table in doc_or_cell.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _fill_paragraph(p, values)
                _walk_tables(cell, values)  # nested tables


def fill_docx_template(template_path: str, values: Dict, output_path: str):
    doc = Document(template_path)
    fill_dynamic_tables(doc, values)
    for p in doc.paragraphs:
        _fill_paragraph(p, values)
    _walk_tables(doc, values)
    for section in doc.sections:
        for p in section.header.paragraphs:
            _fill_paragraph(p, values)
        for p in section.footer.paragraphs:
            _fill_paragraph(p, values)
    doc.save(output_path)


def fill_template(template_path: str, values: Dict, output_path: str):
    """Dispatch based on file extension."""
    if template_path.lower().endswith(".docx"):
        fill_docx_template(template_path, values, output_path)
    else:
        with open(template_path, "r", encoding="utf-8") as f:
            content = f.read()
        filled = fill_text_template(content, values)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(filled)
